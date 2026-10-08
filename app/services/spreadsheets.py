"""Normalized XLSX interchange; PostgreSQL dumps remain the technical backup."""

import asyncio
import hashlib
import json
import logging
import re
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from datetime import datetime, timezone
from enum import Enum
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncEngine, AsyncSession, async_sessionmaker

from app.db.models import (
    AuditLog,
    Club,
    Game,
    GameEvent,
    GamePlayer,
    GameSession,
    Player,
)
from app.domain.enums import (
    EventType,
    GameRole,
    GameStatus,
    RuinReason,
    SessionStatus,
    Team,
    VictoryType,
)
from app.domain.events import validate_event
from app.domain.permissions import can_manage_club
from app.domain.roles import validate_game_roles
from app.domain.victories import validate_victory
from app.services.games import GameError, GameErrorCode, GameService

logger = logging.getLogger(__name__)
COLUMNS = {
    "Players": (
        "player_id",
        "display_name",
        "is_active",
        "telegram_linked",
        "created_at",
        "updated_at",
    ),
    "Sessions": (
        "session_id",
        "started_at",
        "finished_at",
        "status",
        "created_by_player_id",
        "statistician_player_id",
    ),
    "Games": (
        "game_id",
        "session_id",
        "sequence_number",
        "winner_team",
        "victory_type",
        "status",
        "played_at",
        "created_by_player_id",
        "created_at",
        "updated_at",
    ),
    "GamePlayers": ("game_id", "player_id", "player_name", "role"),
    "Events": (
        "event_id",
        "game_id",
        "player_id",
        "player_name",
        "event_type",
        "ruin_reason",
        "self_admitted",
        "comment",
        "created_by_player_id",
        "created_at",
        "updated_at",
    ),
}
MODELS = {
    "Players": Player,
    "Sessions": GameSession,
    "Games": Game,
    "Events": GameEvent,
}
ID_FIELDS = {
    "Players": "player_id",
    "Sessions": "session_id",
    "Games": "game_id",
    "Events": "event_id",
}
REQUIRED_COLUMNS = {
    "Players": {"display_name"},
    "Sessions": {
        "session_id",
        "started_at",
        "created_by_player_id",
        "statistician_player_id",
    },
    "Games": {
        "game_id",
        "session_id",
        "sequence_number",
        "winner_team",
        "played_at",
        "created_by_player_id",
    },
    "GamePlayers": {"game_id", "player_id", "role"},
    "Events": {
        "event_id",
        "game_id",
        "player_id",
        "event_type",
        "created_by_player_id",
    },
}


def workbook_bytes(tables: dict[str, list[dict]]) -> bytes:
    book = Workbook()
    book.remove(book.active)
    try:
        for name, columns in COLUMNS.items():
            sheet = book.create_sheet(name)
            sheet.append(columns)
            for record in tables[name]:
                cells = []
                for column in columns:
                    value = record.get(column)
                    if isinstance(value, (UUID, Enum)):
                        value = str(value) if isinstance(value, UUID) else value.value
                    if isinstance(value, datetime):
                        value = value.astimezone(timezone.utc).isoformat()
                    cell = Cell(sheet, value=value)
                    # Explicit text type preserves =,+,-,@ literally without altering round-trip text.
                    if isinstance(value, str):
                        cell.data_type = "s"
                    cells.append(cell)
                sheet.append(cells)
        stream = BytesIO()
        book.save(stream)
        return stream.getvalue()
    finally:
        book.close()


def export_filename(club_name: str) -> str:
    slug = (
        re.sub(r"[^\w-]+", "_", club_name, flags=re.UNICODE).strip("_-")[:60] or "club"
    )
    return (
        f"secret_hitler_stats_{slug}_{datetime.now(timezone.utc):%Y-%m-%d_%H%M%S}.xlsx"
    )


class ExportService:
    def __init__(self, engine: AsyncEngine) -> None:
        self.engine = engine

    async def export(self, telegram_id: int, club_id: UUID) -> tuple[str, bytes]:
        # Set isolation before beginning any query, including authorization.
        async with self.engine.connect() as connection:
            connection = await connection.execution_options(
                isolation_level="REPEATABLE READ"
            )
            async with connection.begin():
                await connection.execute(text("SET TRANSACTION READ ONLY"))
                async with AsyncSession(bind=connection, expire_on_commit=False) as db:
                    actor = await GameService(async_sessionmaker(self.engine))._actor(
                        db, telegram_id, club_id
                    )
                    if not can_manage_club(actor, club_id):
                        raise GameError(GameErrorCode.NO_ACCESS)
                    club = await db.get(Club, club_id)
                    players = (
                        await db.scalars(
                            select(Player)
                            .where(Player.club_id == club_id)
                            .order_by(Player.id)
                        )
                    ).all()
                    sessions = (
                        await db.scalars(
                            select(GameSession)
                            .where(GameSession.club_id == club_id)
                            .order_by(GameSession.id)
                        )
                    ).all()
                    game_query = (
                        select(Game.id)
                        .join(GameSession)
                        .where(GameSession.club_id == club_id)
                    )
                    games = (
                        await db.scalars(
                            select(Game)
                            .where(Game.id.in_(game_query))
                            .order_by(Game.id)
                        )
                    ).all()
                    roles = (
                        await db.scalars(
                            select(GamePlayer)
                            .where(GamePlayer.game_id.in_(game_query))
                            .order_by(GamePlayer.game_id, GamePlayer.player_id)
                        )
                    ).all()
                    events = (
                        await db.scalars(
                            select(GameEvent)
                            .where(GameEvent.game_id.in_(game_query))
                            .order_by(GameEvent.id)
                        )
                    ).all()
                    names = {p.id: p.display_name for p in players}
                    tables = {}
                    for sheet, records in (
                        ("Players", players),
                        ("Sessions", sessions),
                        ("Games", games),
                        ("GamePlayers", roles),
                        ("Events", events),
                    ):
                        tables[sheet] = []
                        for row in records:
                            record = {}
                            for column in COLUMNS[sheet]:
                                if column == ID_FIELDS.get(sheet):
                                    record[column] = row.id
                                elif column == "player_name":
                                    record[column] = names[row.player_id]
                                elif column == "telegram_linked":
                                    record[column] = row.telegram_user_id is not None
                                else:
                                    record[column] = getattr(row, column)
                            tables[sheet].append(record)
                    filename = export_filename(club.name)
        payload = await asyncio.to_thread(workbook_bytes, tables)
        logger.info("Club exported: club=%s games=%s", club_id, len(games))
        return filename, payload


@dataclass
class ImportReport:
    tables: dict[str, list[dict]] = field(
        default_factory=lambda: {s: [] for s in COLUMNS}
    )
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    fingerprint: str = ""

    @property
    def counts(self) -> dict[str, int]:
        return {name.lower(): len(rows) for name, rows in self.tables.items()}


def _bool(value: object, default: bool = False) -> bool:
    if value is None:
        return default
    if type(value) is bool or type(value) is int and value in (0, 1):
        return bool(value)
    if isinstance(value, str) and value.upper() in ("TRUE", "FALSE"):
        return value.upper() == "TRUE"
    raise ValueError("Expected TRUE/FALSE")


def _date(value: str | datetime | None) -> datetime | None:
    if value is None:
        return None
    result = (
        value if isinstance(value, datetime) else datetime.fromisoformat(str(value))
    )
    if result.tzinfo is None:
        raise ValueError("Timestamp must include timezone (ISO 8601)")
    return result


def read_import(source: Path | bytes) -> ImportReport:
    report = ImportReport()
    try:
        if isinstance(source, Path):
            if not source.is_file() or source.suffix.lower() != ".xlsx":
                raise ValueError("Input must be an existing .xlsx file")
            if source.stat().st_size > 20 * 1024 * 1024:
                raise ValueError("XLSX exceeds 20 MiB")
            data = source.read_bytes()
        else:
            data = source
        if len(data) > 20 * 1024 * 1024:
            raise ValueError("XLSX exceeds 20 MiB")
        with ZipFile(BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Expanded XLSX exceeds 100 MiB")
        book = load_workbook(
            BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            for name in COLUMNS:
                if name not in book.sheetnames:
                    report.errors.append(f"Missing sheet: {name}")
                    continue
                sheet = book[name]
                if (sheet.max_row or 0) > 100001 or (sheet.max_column or 0) > 64:
                    report.errors.append(
                        f"{name}: sheet dimensions exceed 100000 data rows / 64 columns"
                    )
                    continue
                iterator = sheet.iter_rows()
                first = next(iterator, ())
                headers = [c.value for c in first]
                if len(set(headers)) != len(headers) or any(
                    not isinstance(h, str) for h in headers
                ):
                    report.errors.append(f"{name}: invalid/duplicate headers")
                    continue
                missing = REQUIRED_COLUMNS[name] - set(headers)
                if missing:
                    report.errors.append(f"{name}: missing columns {sorted(missing)}")
                    continue
                for number, cells in enumerate(iterator, 2):
                    if number > 100001:
                        raise ValueError("Sheet exceeds 100000 rows")
                    if all(c.value is None for c in cells):
                        continue
                    if any(c.data_type in ("f", "e") for c in cells):
                        report.errors.append(
                            f"{name} row {number}: formulas/errors are forbidden"
                        )
                        continue
                    row = dict(zip(headers, (c.value for c in cells)))
                    row["_row"] = number
                    report.tables[name].append(row)
        finally:
            book.close()
        # Identity includes normalized source contents, independent of XLSX ZIP metadata.
        report.fingerprint = hashlib.sha256(
            json.dumps(
                report.tables, sort_keys=True, default=str, ensure_ascii=False
            ).encode()
        ).hexdigest()
        _validate_import(report)
    except ValueError as error:
        report.errors.append(f"Invalid workbook: {error}")
    except Exception as error:
        report.errors.append(f"Invalid workbook: {type(error).__name__}")
    return report


def _validate_import(report: ImportReport) -> None:
    tables = report.tables
    players, sessions, games = {}, {}, {}
    names = {}
    original_ids = {}

    def check(name, row, operation):
        try:
            operation(row)
        except (ValueError, KeyError, TypeError) as error:
            report.errors.append(f"{name} row {row['_row']}: {error}")

    def player_ref(value):
        try:
            result = UUID(str(value))
        except ValueError:
            result = None
        if result in players:
            return result
        if value in names and len(names[value]) == 1:
            return names[value][0]
        raise ValueError("Unknown or ambiguous player reference")

    def player(row):
        name = row["display_name"]
        if not isinstance(name, str) or not name.strip() or len(name) > 255:
            raise ValueError("Invalid display_name")
        given = row.get("player_id")
        pid = UUID(str(given)) if given else uuid4()
        if pid in players:
            raise ValueError("Duplicate player_id")
        original_ids[pid] = bool(given)
        row.update(
            id=pid,
            display_name=name.strip(),
            is_active=_bool(row.get("is_active"), True),
        )
        players[pid] = row
        names.setdefault(name.strip(), []).append(pid)

    for row in tables["Players"]:
        check("Players", row, player)
    folded_names = {}
    for name, ids in names.items():
        folded_names.setdefault(name.casefold(), []).extend(ids)
    for name, ids in folded_names.items():
        if len(ids) > 1:
            message = f"Ambiguous player name {name!r}; rows {[players[i]['_row'] for i in ids]}"
            (
                report.warnings if all(original_ids[i] for i in ids) else report.errors
            ).append(message)

    # Compare lexical neighbours: bounded work even for a large input, no auto mapping.
    ordered_names = sorted(folded_names)
    for left, right in zip(ordered_names, ordered_names[1:]):
        if SequenceMatcher(None, left, right).ratio() >= 0.8 or (
            min(len(left), len(right)) >= 3
            and (left.startswith(right) or right.startswith(left))
        ):
            report.warnings.append(
                f"Similar names {left!r} / {right!r}; check explicit Player mapping"
            )

    def session(row):
        sid = UUID(str(row["session_id"]))
        if sid in sessions:
            raise ValueError("Duplicate session_id")
        row.update(
            id=sid,
            status=SessionStatus(row.get("status") or "FINISHED"),
            started_at=_date(row["started_at"]),
            finished_at=_date(row.get("finished_at")),
        )
        if row["started_at"] is None:
            raise ValueError("started_at required")
        for key in ("created_by_player_id", "statistician_player_id"):
            row[key] = player_ref(row[key])
        sessions[sid] = row

    for row in tables["Sessions"]:
        check("Sessions", row, session)
    if sum(s["status"] is SessionStatus.ACTIVE for s in sessions.values()) > 1:
        report.errors.append("Only one ACTIVE session is allowed")

    def game(row):
        gid, sid = UUID(str(row["game_id"])), UUID(str(row["session_id"]))
        if gid in games:
            raise ValueError("Duplicate Game ID")
        if sid not in sessions:
            raise ValueError("Unknown session reference")
        seq = row["sequence_number"]
        if type(seq) is not int or seq < 1:
            raise ValueError("Invalid sequence_number")
        if any(
            g["session_id"] == sid and g["sequence_number"] == seq
            for g in games.values()
        ):
            raise ValueError("Duplicate session sequence_number")
        row.update(
            id=gid,
            session_id=sid,
            winner_team=Team(row["winner_team"]),
            victory_type=VictoryType(row.get("victory_type") or "UNKNOWN"),
            status=GameStatus(row.get("status") or "ACTIVE"),
            played_at=_date(row["played_at"]),
            created_by_player_id=player_ref(row["created_by_player_id"]),
        )
        if row["played_at"] is None:
            raise ValueError("played_at required")
        validate_victory(row["winner_team"], row["victory_type"], allow_unknown=True)
        games[gid] = row

    for row in tables["Games"]:
        check("Games", row, game)
    assignments = {gid: [] for gid in games}

    def role(row):
        gid = UUID(str(row["game_id"]))
        if gid not in games:
            raise ValueError("Unknown game reference")
        row.update(
            game_id=gid,
            player_id=player_ref(row["player_id"]),
            role=GameRole(row["role"]),
        )
        assignments[gid].append((row["player_id"], row["role"]))

    for row in tables["GamePlayers"]:
        check("GamePlayers", row, role)
    for gid, roles in assignments.items():
        check(
            "Games",
            games[gid],
            lambda row: validate_game_roles(
                [pid for pid, _ in roles], roles, club_player_ids=players
            ),
        )
    event_ids = set()

    def event(row):
        eid, gid = UUID(str(row["event_id"])), UUID(str(row["game_id"]))
        if eid in event_ids:
            raise ValueError("Duplicate event_id")
        if gid not in games:
            raise ValueError("Unknown game reference")
        pid = player_ref(row["player_id"])
        fields = validate_event(
            pid,
            [p for p, _ in assignments[gid]],
            EventType(row["event_type"]),
            RuinReason(row["ruin_reason"]) if row.get("ruin_reason") else None,
            _bool(row.get("self_admitted")),
            row.get("comment"),
        )
        row.update(
            id=eid,
            game_id=gid,
            player_id=pid,
            created_by_player_id=player_ref(row["created_by_player_id"]),
            **vars(fields),
        )
        event_ids.add(eid)

    for row in tables["Events"]:
        check("Events", row, event)
    for sheet in ("Players", "Games", "Events"):
        for row in tables[sheet]:
            for key in ("created_at", "updated_at"):
                if row.get(key) is not None:
                    check(sheet, row, lambda r, k=key: r.update({k: _date(r[k])}))


class ImportService:
    """CLI operator access; never exposed as an unauthenticated Telegram handler."""

    def __init__(self, factory: async_sessionmaker[AsyncSession]) -> None:
        self.factory = factory

    async def run(
        self, club_id: UUID, report: ImportReport, *, apply: bool = False
    ) -> ImportReport:
        # Do not mutate reusable report errors or records during DB validation.
        report = ImportReport(
            report.tables,
            list(report.errors),
            list(report.warnings),
            report.fingerprint,
        )
        async with self.factory() as db, db.begin():
            club = await db.scalar(
                select(Club).where(Club.id == club_id).with_for_update()
                if apply
                else select(Club).where(Club.id == club_id)
            )
            if club is None:
                report.errors.append("Target club does not exist")
            prior = await db.scalar(
                select(AuditLog.id).where(
                    AuditLog.club_id == club_id,
                    AuditLog.action == "IMPORT",
                    AuditLog.after_data["fingerprint"].astext == report.fingerprint,
                )
            )
            if prior:
                report.errors.append("Dataset already imported")
            existing_players = set()
            for sheet, model in MODELS.items():
                ids = [r["id"] for r in report.tables[sheet] if "id" in r]
                existing = (
                    await db.scalars(select(model).where(model.id.in_(ids)))
                ).all()
                for obj in existing:
                    if sheet == "Players" and obj.club_id == club_id:
                        row = next(
                            r for r in report.tables[sheet] if r.get("id") == obj.id
                        )
                        if (
                            row["display_name"] == obj.display_name
                            and row["is_active"] == obj.is_active
                        ):
                            existing_players.add(obj.id)
                            report.warnings.append(
                                f"Existing Player {obj.id} reused without update"
                            )
                            continue
                    report.errors.append(f"{sheet}: duplicate/conflicting ID {obj.id}")
            known = (
                await db.execute(
                    select(Player.id, Player.display_name).where(
                        Player.club_id == club_id
                    )
                )
            ).all()
            for row in report.tables["Players"]:
                if row.get("id") not in existing_players and any(
                    name.casefold() == str(row.get("display_name", "")).casefold()
                    for _, name in known
                ):
                    report.errors.append(
                        f"Players row {row['_row']}: existing name needs explicit Player UUID mapping"
                    )
            if any(
                r.get("status") is SessionStatus.ACTIVE
                for r in report.tables["Sessions"]
            ):
                if await db.scalar(
                    select(GameSession.id).where(
                        GameSession.club_id == club_id,
                        GameSession.status == SessionStatus.ACTIVE,
                    )
                ):
                    report.errors.append("Target club already has an ACTIVE session")
            if report.errors or not apply:
                return report

            def values(row, allowed):
                return {k: v for k, v in row.items() if k in allowed}

            for row in report.tables["Players"]:
                if row["id"] not in existing_players:
                    db.add(
                        Player(
                            club_id=club_id,
                            telegram_user_id=None,
                            **values(
                                row,
                                (
                                    "id",
                                    "display_name",
                                    "is_active",
                                    "created_at",
                                    "updated_at",
                                ),
                            ),
                        )
                    )
            await db.flush()
            for row in report.tables["Sessions"]:
                db.add(
                    GameSession(
                        club_id=club_id,
                        **values(
                            row,
                            (
                                "id",
                                "started_at",
                                "finished_at",
                                "status",
                                "created_by_player_id",
                                "statistician_player_id",
                            ),
                        ),
                    )
                )
            await db.flush()
            for row in report.tables["Games"]:
                db.add(
                    Game(
                        **values(
                            row,
                            (
                                "id",
                                "session_id",
                                "sequence_number",
                                "winner_team",
                                "victory_type",
                                "status",
                                "played_at",
                                "created_by_player_id",
                                "created_at",
                                "updated_at",
                            ),
                        )
                    )
                )
            await db.flush()
            for row in report.tables["GamePlayers"]:
                db.add(GamePlayer(**values(row, ("game_id", "player_id", "role"))))
            await db.flush()
            for row in report.tables["Events"]:
                db.add(
                    GameEvent(
                        **values(
                            row,
                            (
                                "id",
                                "game_id",
                                "player_id",
                                "event_type",
                                "ruin_reason",
                                "self_admitted",
                                "comment",
                                "created_by_player_id",
                                "created_at",
                                "updated_at",
                            ),
                        )
                    )
                )
            db.add(
                AuditLog(
                    club_id=club_id,
                    actor_player_id=None,
                    entity_type="Club",
                    entity_id=club_id,
                    action="IMPORT",
                    after_data={**report.counts, "fingerprint": report.fingerprint},
                )
            )
            await db.flush()
        logger.info(
            "Legacy import complete: club=%s summary=%s", club_id, report.counts
        )
        return report
