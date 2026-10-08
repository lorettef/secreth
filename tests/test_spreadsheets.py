"""XLSX domain/security and transactional, scoped PostgreSQL interchange."""

import asyncio
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.db.base import Base
from app.db.models import (
    AuditLog,
    Club,
    Game,
    GameEvent,
    GamePlayer,
    GameSession,
    Player,
)
from app.domain.enums import (
    EventType,
    GameStatus,
    RuinReason,
    SessionStatus,
    Team,
    VictoryType,
)
from app.services.events import EventService
from app.services.games import GameError, GameService
from app.services.spreadsheets import (
    COLUMNS,
    ExportService,
    ImportService,
    export_filename,
    read_import,
    workbook_bytes,
)
from test_games_integration import seed, test_database_url  # noqa: F401


@pytest_asyncio.fixture
async def exchange(test_database_url):
    engine = create_async_engine(test_database_url, hide_parameters=True)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    club, evening, players, _ = await seed(factory)
    other, _, _, _ = await seed(factory, telegram_ids=(8901, 8902))
    games, events = GameService(factory), EventService(factory)
    game, _ = await games.create_game(
        8801,
        evening,
        uuid4(),
        players[:5],
        players[0],
        (players[1],),
        Team.FASCIST,
        VictoryType.FASCIST_POLICIES,
    )
    await events.create_event(
        8801,
        game.id,
        uuid4(),
        players[1],
        EventType.RUIN,
        RuinReason.WRONG_VOTE,
        True,
        "Я ошибся — <b>&_*[=Юникод",
    )
    await events.create_event(
        8801, game.id, uuid4(), players[2], EventType.IQ_200, comment="=SUM(1,2)"
    )
    game2, _ = await games.create_game(
        8801,
        evening,
        uuid4(),
        players[:6],
        players[0],
        (players[1],),
        Team.LIBERAL,
        VictoryType.HITLER_KILLED,
    )
    async with factory() as db, db.begin():
        (await db.get(Player, players[5])).is_active = False
        (await db.get(Game, game2.id)).status = GameStatus.VOID
        (await db.get(GameSession, evening)).status = SessionStatus.FINISHED
    yield engine, factory, club, evening, players, game.id
    # The entire dedicated *_test DB is disposable; no other test is running in parallel.
    async with factory() as db, db.begin():
        for table in reversed(Base.metadata.sorted_tables):
            await db.execute(delete(table))
    await engine.dispose()


def raw_tables(payload):
    book = load_workbook(BytesIO(payload), read_only=True, data_only=False)
    try:
        tables = {}
        for sheet in book:
            rows = list(sheet.values)
            tables[sheet.title] = [dict(zip(rows[0], row)) for row in rows[1:]]
        return tables
    finally:
        book.close()


@pytest.mark.asyncio
async def test_export_all_sheets_scoped_unicode_ids_void_inactive_and_no_disk(
    exchange, tmp_path, monkeypatch
):
    engine, factory, club, evening, players, game = exchange
    monkeypatch.chdir(tmp_path)
    filename, payload = await ExportService(engine).export(8802, club)
    assert filename.endswith(".xlsx") and list(tmp_path.iterdir()) == []
    tables = raw_tables(payload)
    assert set(tables) == set(COLUMNS)
    assert len(tables["Players"]) == 10 and len(tables["Games"]) == 2
    assert str(game) in {g["game_id"] for g in tables["Games"]}
    assert {g["status"] for g in tables["Games"]} == {"VOID", "ACTIVE"}
    assert {p["role"] for p in tables["GamePlayers"]} == {
        "LIBERAL",
        "FASCIST",
        "HITLER",
    }
    assert not next(p for p in tables["Players"] if p["player_id"] == str(players[5]))[
        "is_active"
    ]
    assert not next(p for p in tables["Players"] if p["player_id"] == str(players[8]))[
        "telegram_linked"
    ]
    assert (
        "Я ошибся" in tables["Events"][0]["comment"]
        or "Я ошибся" in tables["Events"][1]["comment"]
    )
    assert read_import(payload).errors == []


@pytest.mark.asyncio
@pytest.mark.parametrize("actor", [8801, 8902, 9999])
async def test_export_permissions_and_other_club(exchange, actor):
    with pytest.raises(GameError):
        await ExportService(exchange[0]).export(actor, exchange[2])


@pytest.mark.parametrize(
    "value", ["=1+1", "+CMD", "-CMD", "@SUM", "<tg-spoiler><b>&_*[Юникод"]
)
def test_formula_injection_is_literal_and_roundtrip(value):
    tables = {sheet: [] for sheet in COLUMNS}
    tables["Players"] = [{"display_name": value}]
    payload = workbook_bytes(tables)
    book = load_workbook(BytesIO(payload), data_only=False)
    assert book["Players"]["B2"].data_type == "s"
    assert book["Players"]["B2"].value == value
    book.close()


def test_import_rejects_formulas():
    tables = {sheet: [] for sheet in COLUMNS}
    book = load_workbook(BytesIO(workbook_bytes(tables)))
    book["Players"].append([None, "=1+1"])
    output = BytesIO()
    book.save(output)
    book.close()
    assert any("formulas" in error for error in read_import(output.getvalue()).errors)


@pytest.mark.parametrize("name", ["../../club", "/tmp/file", "😈", "x" * 1000])
def test_filename_cannot_escape_directory(name):
    filename = export_filename(name)
    assert Path(filename).name == filename and len(filename) < 120


@pytest.mark.parametrize("path", ["missing.xlsx", ".", "wrong.csv"])
def test_import_invalid_paths(tmp_path, path):
    assert read_import(tmp_path / path).errors


async def prepared(exchange):
    engine, factory, club, _, _, _ = exchange
    _, payload = await ExportService(engine).export(8802, club)
    tables = raw_tables(payload)
    # A second club in the same DB needs new globally unique entity IDs.
    mapping = {
        row[key]: str(uuid4())
        for sheet, key in [
            ("Players", "player_id"),
            ("Sessions", "session_id"),
            ("Games", "game_id"),
            ("Events", "event_id"),
        ]
        for row in tables[sheet]
    }
    for rows in tables.values():
        for row in rows:
            for key, value in list(row.items()):
                if isinstance(value, str) and value in mapping:
                    row[key] = mapping[value]
    async with factory() as db, db.begin():
        target = Club(name="Import target", timezone="UTC")
        db.add(target)
        await db.flush()
        target_id = target.id
    return target_id, tables


@pytest.mark.asyncio
async def test_import_dry_run_apply_unknown_stats_audit_and_duplicate_dataset(exchange):
    target, tables = await prepared(exchange)
    tables["Games"][0]["victory_type"] = "UNKNOWN"
    report = read_import(workbook_bytes(tables))
    assert not report.errors
    service = ImportService(exchange[1])
    dry = await service.run(target, report)
    assert not dry.errors
    async with exchange[1]() as db:
        assert (
            await db.scalar(
                select(func.count(Player.id)).where(Player.club_id == target)
            )
            == 0
        )
        assert (
            await db.scalar(
                select(func.count(AuditLog.id)).where(AuditLog.club_id == target)
            )
            == 0
        )
    applied = await service.run(target, report, apply=True)
    assert not applied.errors
    async with exchange[1]() as db:
        assert (
            await db.scalar(
                select(func.count(Player.id)).where(Player.club_id == target)
            )
            == 10
        )
        assert (
            await db.scalar(
                select(func.count(Player.id)).where(
                    Player.club_id == target, Player.telegram_user_id.is_not(None)
                )
            )
            == 0
        )
        audit = await db.scalar(select(AuditLog).where(AuditLog.club_id == target))
        assert (
            audit.action == "IMPORT"
            and audit.actor_player_id is None
            and audit.after_data["games"] == 2
        )
    second = await service.run(target, read_import(workbook_bytes(tables)), apply=True)
    assert second.errors and any("already imported" in e for e in second.errors)


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "defect",
    [
        "role",
        "duplicate_name",
        "unknown_ref",
        "duplicate_game",
        "winner",
        "missing_reason",
        "comment",
        "date",
        "sequence",
        "foreign_id",
    ],
)
async def test_invalid_import_is_atomic_and_reports_rows(exchange, defect):
    target, tables = await prepared(exchange)
    if defect == "role":
        tables["GamePlayers"][0]["role"] = "INVALID"
    if defect == "duplicate_name":
        tables["Players"][0]["player_id"] = None
        tables["Players"][1]["player_id"] = None
        tables["Players"][1]["display_name"] = tables["Players"][0]["display_name"]
    if defect == "unknown_ref":
        tables["GamePlayers"][0]["player_id"] = str(uuid4())
    if defect == "duplicate_game":
        tables["Games"].append(dict(tables["Games"][0]))
    if defect == "winner":
        tables["Games"][0].update(winner_team="HITLER")
    if defect == "missing_reason":
        for e in tables["Events"]:
            e.update(event_type="RUIN", ruin_reason=None)
    if defect == "comment":
        tables["Events"][0]["comment"] = "x" * 501
    if defect == "date":
        tables["Games"][0]["played_at"] = "2020-01-01T00:00:00"
    if defect == "sequence":
        tables["Games"][0]["sequence_number"] = -1
    if defect == "foreign_id":
        tables["Players"][0]["player_id"] = str(exchange[4][0])
    result = await ImportService(exchange[1]).run(
        target, read_import(workbook_bytes(tables)), apply=True
    )
    assert result.errors, defect
    async with exchange[1]() as db:
        assert (
            await db.scalar(
                select(func.count(Player.id)).where(Player.club_id == target)
            )
            == 0
        )
        assert (
            await db.scalar(
                select(func.count(GameSession.id)).where(GameSession.club_id == target)
            )
            == 0
        )
        assert (
            await db.scalar(
                select(func.count(AuditLog.id)).where(AuditLog.club_id == target)
            )
            == 0
        )


@pytest.mark.asyncio
async def test_import_missing_player_uuid_exact_name_references_and_duplicate_fingerprint(
    exchange,
):
    target, tables = await prepared(exchange)
    for row in tables["Players"]:
        original, name = row["player_id"], row["display_name"]
        row["player_id"] = None
        for sheet in ("Sessions", "Games", "GamePlayers", "Events"):
            for item in tables[sheet]:
                for key in (
                    "player_id",
                    "created_by_player_id",
                    "statistician_player_id",
                ):
                    if item.get(key) == original:
                        item[key] = name
    result = await ImportService(exchange[1]).run(
        target, read_import(workbook_bytes(tables)), apply=True
    )
    assert not result.errors
    second = await ImportService(exchange[1]).run(
        target, read_import(workbook_bytes(tables)), apply=True
    )
    assert any("already imported" in error for error in second.errors)


@pytest.mark.asyncio
async def test_export_import_same_uuids_roundtrip_clean_target(exchange):
    engine, factory, club, _, _, _ = exchange
    _, payload = await ExportService(engine).export(8802, club)
    before = raw_tables(payload)
    # Simulate recovery into a fresh club after removing the disposable source dataset.
    async with factory() as db, db.begin():
        for table in reversed(Base.metadata.sorted_tables):
            await db.execute(delete(table))
        target = Club(name="Restored", timezone="UTC")
        db.add(target)
        await db.flush()
        target_id = target.id
    result = await ImportService(factory).run(
        target_id, read_import(payload), apply=True
    )
    assert not result.errors
    async with factory() as db:
        for sheet, model in [
            ("Players", Player),
            ("Sessions", GameSession),
            ("Games", Game),
            ("Events", GameEvent),
        ]:
            field = {
                "Players": "player_id",
                "Sessions": "session_id",
                "Games": "game_id",
                "Events": "event_id",
            }[sheet]
            assert {str(i) for i in (await db.scalars(select(model.id))).all()} == {
                r[field] for r in before[sheet]
            }
        roles = (
            await db.execute(
                select(GamePlayer.game_id, GamePlayer.player_id, GamePlayer.role)
            )
        ).all()
        assert {(str(g), str(p), r.value) for g, p, r in roles} == {
            (r["game_id"], r["player_id"], r["role"]) for r in before["GamePlayers"]
        }
        events = (await db.scalars(select(GameEvent))).all()
        assert {
            (str(e.id), e.event_type.value, e.comment, e.self_admitted) for e in events
        } == {
            (r["event_id"], r["event_type"], r["comment"], r["self_admitted"])
            for r in before["Events"]
        }
        games = (await db.scalars(select(Game))).all()
        assert {
            (str(g.id), g.winner_team.value, g.victory_type.value, g.status.value)
            for g in games
        } == {
            (r["game_id"], r["winner_team"], r["victory_type"], r["status"])
            for r in before["Games"]
        }


@pytest.mark.asyncio
async def test_export_snapshot_survives_concurrent_event_write(exchange, monkeypatch):
    from sqlalchemy.ext.asyncio import AsyncSession

    engine, factory, club, _, players, game = exchange
    read_games, event_committed = asyncio.Event(), asyncio.Event()
    original = AsyncSession.scalars

    async def paused(db, statement, *args, **kwargs):
        result = await original(db, statement, *args, **kwargs)
        if statement.column_descriptions[0].get("entity") is Game:
            read_games.set()
            await event_committed.wait()
        return result

    monkeypatch.setattr(AsyncSession, "scalars", paused)
    task = asyncio.create_task(ExportService(engine).export(8802, club))
    try:
        await asyncio.wait_for(read_games.wait(), 10)
        await EventService(factory).create_event(
            8802, game, uuid4(), players[0], EventType.PALIVO
        )
    finally:
        event_committed.set()
    _, payload = await task
    assert len(raw_tables(payload)["Events"]) == 2
    async with factory() as db:
        assert await db.scalar(select(func.count(GameEvent.id))) == 3


@pytest.mark.asyncio
async def test_import_transaction_rolls_back_on_database_write_failure(exchange):
    from sqlalchemy import event

    target, tables = await prepared(exchange)
    report = read_import(workbook_bytes(tables))
    assert not report.errors

    def fail(conn, cursor, statement, parameters, context, executemany):
        if statement.startswith("INSERT INTO game_events"):
            raise RuntimeError("simulated storage failure")

    event.listen(exchange[0].sync_engine, "before_cursor_execute", fail)
    try:
        with pytest.raises(RuntimeError, match="storage"):
            await ImportService(exchange[1]).run(target, report, apply=True)
    finally:
        event.remove(exchange[0].sync_engine, "before_cursor_execute", fail)
    async with exchange[1]() as db:
        assert (
            await db.scalar(
                select(func.count(Player.id)).where(Player.club_id == target)
            )
            == 0
        )
        assert (
            await db.scalar(
                select(func.count(GameSession.id)).where(GameSession.club_id == target)
            )
            == 0
        )
        assert (
            await db.scalar(
                select(func.count(AuditLog.id)).where(AuditLog.club_id == target)
            )
            == 0
        )


@pytest.mark.asyncio
async def test_concurrent_import_serializes_same_dataset(exchange):
    target, tables = await prepared(exchange)
    payload = workbook_bytes(tables)
    results = await asyncio.gather(
        *(
            ImportService(exchange[1]).run(target, read_import(payload), apply=True)
            for _ in range(2)
        )
    )
    assert sum(not r.errors for r in results) == 1
    async with exchange[1]() as db:
        assert (
            await db.scalar(
                select(func.count(Player.id)).where(Player.club_id == target)
            )
            == 10
        )
        assert (
            await db.scalar(
                select(func.count(AuditLog.id)).where(
                    AuditLog.club_id == target, AuditLog.action == "IMPORT"
                )
            )
            == 1
        )


@pytest.mark.asyncio
async def test_import_stats_and_historical_players_available_without_telegram(exchange):
    from app.db.models import ClubMembership, TelegramUser
    from app.domain.enums import MembershipRole, MembershipStatus
    from app.services.stats import StatsService

    target, tables = await prepared(exchange)
    result = await ImportService(exchange[1]).run(
        target, read_import(workbook_bytes(tables)), apply=True
    )
    assert not result.errors
    async with exchange[1]() as db, db.begin():
        user = await db.scalar(
            select(TelegramUser.id).where(TelegramUser.telegram_user_id == 8802)
        )
        db.add(
            ClubMembership(
                club_id=target,
                telegram_user_id=user,
                role=MembershipRole.ADMIN,
                status=MembershipStatus.ACTIVE,
            )
        )
    stats = await StatsService(exchange[1]).get_club_stats(8802, target)
    assert stats.score.games_count == 1 and stats.score.fascist_wins == 1
    assert (
        stats.event_counts[EventType.RUIN] == 1
        and stats.event_counts[EventType.IQ_200] == 1
    )
    async with exchange[1]() as db:
        assert (
            await db.scalar(
                select(func.count(Player.id)).where(
                    Player.club_id == target, Player.telegram_user_id.is_not(None)
                )
            )
            == 0
        )


def test_import_checks_mandatory_headers_even_on_empty_sheet():
    book = load_workbook(BytesIO(workbook_bytes({s: [] for s in COLUMNS})))
    book["Games"].delete_cols(1)
    output = BytesIO()
    book.save(output)
    book.close()
    assert any(
        "Games: missing columns" in error
        for error in read_import(output.getvalue()).errors
    )


@pytest.mark.asyncio
async def test_real_cli_dry_run_apply_and_duplicate_rejection(exchange, tmp_path):
    import os
    import subprocess
    import sys

    target, tables = await prepared(exchange)
    filename = tmp_path / "история.xlsx"
    filename.write_bytes(workbook_bytes(tables))
    environment = dict(os.environ, DATABASE_URL=os.environ["TEST_DATABASE_URL"])

    def run(mode):
        return subprocess.run(
            [
                sys.executable,
                "scripts/import_legacy.py",
                str(filename),
                mode,
                "--club-id",
                str(target),
            ],
            env=environment,
            capture_output=True,
            text=True,
        )

    dry = await asyncio.to_thread(run, "--dry-run")
    assert dry.returncode == 0 and "APPLY NOT PERFORMED" in dry.stdout
    async with exchange[1]() as db:
        assert (
            await db.scalar(
                select(func.count(Player.id)).where(Player.club_id == target)
            )
            == 0
        )
    applied = await asyncio.to_thread(run, "--apply")
    assert applied.returncode == 0 and "APPLIED" in applied.stdout
    duplicate = await asyncio.to_thread(run, "--apply")
    assert duplicate.returncode == 2 and "APPLY NOT PERFORMED" in duplicate.stdout


def test_import_rejects_huge_sparse_sheet_dimensions():
    book = load_workbook(BytesIO(workbook_bytes({s: [] for s in COLUMNS})))
    book["Players"].cell(row=100002, column=1, value="end")
    output = BytesIO()
    book.save(output)
    book.close()
    assert any("dimensions" in e for e in read_import(output.getvalue()).errors)


@pytest.mark.parametrize(
    "names, expected",
    [(["Макс", "макс"], "Ambiguous"), (["Макс", "Максим"], "Similar")],
)
def test_import_reports_case_ambiguity_and_similar_names_without_merging(
    names, expected
):
    tables = {s: [] for s in COLUMNS}
    tables["Players"] = [{"display_name": name} for name in names]
    report = read_import(workbook_bytes(tables))
    assert any(expected in item for item in report.errors + report.warnings)
    assert len({r["id"] for r in report.tables["Players"]}) == 2
